"""Verify database access and export output without changing production records."""

import argparse
import csv
import json
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import urlsplit
from urllib.request import urlopen

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import select, text  # noqa: E402

import app.models  # noqa: E402, F401
from app.db.session import SessionLocal  # noqa: E402
from app.models.exports import Export  # noqa: E402
from app.models.website import Website  # noqa: E402


def verify_api() -> None:
    from app.services.system_status import build_queue_status

    with SessionLocal() as db:
        db.execute(text("SET TRANSACTION READ ONLY"))
        assert db.scalar(text("SELECT version_num FROM alembic_version")) == "0066"
        for table in ("security_incidents", "queue_dead_letters"):
            assert db.scalar(text("SELECT to_regclass(:name)"), {"name": f"public.{table}"})
            assert db.scalar(
                text("SELECT has_table_privilege(current_user, :name, 'SELECT')"),
                {"name": f"public.{table}"},
            ), f"API cannot read {table}"
            db.execute(text(f"SELECT 1 FROM public.{table} LIMIT 1"))
    with urlopen("http://localhost:8000/health", timeout=10) as response:
        assert response.status == 200
    status = build_queue_status()
    print(json.dumps(status, indent=2))
    for name, queue in status["queues"].items():
        assert queue["workers"] > 0, f"No registered worker for {name}"
    print("API, migration 0066, monitoring tables and worker registrations verified.")


def verify_exports() -> None:
    from app.services.exports import EXPORT_ROOT, _write_csv, _write_excel

    with SessionLocal() as db:
        db.execute(text("SET TRANSACTION READ ONLY"))
        assert db.scalar(text("SELECT current_user")) == "seo_export"
        for column in ("id", "display_name", "email", "password_hash", "mfa_secret_encrypted"):
            allowed = db.scalar(
                text(
                    "SELECT has_column_privilege(current_user, 'public.users', :column, 'SELECT')"
                ),
                {"column": column},
            )
            assert allowed == (column in {"id", "display_name", "email"}), column
        # PostgreSQL checks privileges even for an empty result.
        db.execute(text("SELECT id, display_name, email FROM public.users LIMIT 0"))
        sites = [
            site
            for site in db.scalars(select(Website))
            if (urlsplit(site.base_url).hostname or "").lower().removeprefix("www.")
            == "schipperkozijnen.nl"
        ]
        assert len(sites) == 1, (
            "Expected exactly one Schipper website; stop and inspect configuration"
        )
        site = sites[0]
        with TemporaryDirectory(prefix="operations-check-", dir=EXPORT_ROOT) as directory:
            path = Path(directory)
            _write_csv(db, Export(website_id=site.id, export_type="urls"), path / "urls.csv")
            with (path / "urls.csv").open(newline="", encoding="utf-8") as handle:
                rows = list(csv.reader(handle))
            assert len(rows) > 1, "CSV has no URL data"
            _write_excel(db, site.id, path / "report.xlsx")
            workbook = load_workbook(path / "report.xlsx", read_only=True)
            try:
                assert {"Metadata", "Urls", "Issues", "Tasks"} <= set(workbook.sheetnames)
                assert workbook["Urls"].max_row == len(rows)
                print(
                    f"CSV and Excel verified: {len(rows) - 1} URLs; "
                    f"{len(workbook.sheetnames)} sheets."
                )
            finally:
                workbook.close()
    print("Export role verified; temporary test files removed; database unchanged.")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=("api", "exports"))
    args = parser.parse_args()
    try:
        (verify_api if args.mode == "api" else verify_exports)()
    except Exception as error:
        # Avoid printing SQL parameters, user data or connection strings.
        detail = str(error) if isinstance(error, AssertionError) else type(error).__name__
        print(f"Verification failed: {detail}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
