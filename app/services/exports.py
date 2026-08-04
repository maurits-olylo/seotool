import csv
import json
import uuid
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.session import SessionLocal
from app.models.crawl import CrawlRun, UrlLink
from app.models.discovery import Url, UrlSource
from app.models.exports import Export
from app.models.issues import Change, Issue
from app.models.jobs import JobListing
from app.models.recommendations import RecommendationTask
from app.models.user import User
from app.models.website import Website
from app.services.job_posting import ACTIVE_JOB_ISSUE_STATUSES, JOB_ISSUE_TYPES

EXPORT_ROOT = Path("/app/exports")


def generate_export(export_id: str) -> None:
    with SessionLocal() as db:
        export = db.get(Export, uuid.UUID(export_id))
        if export is None or export.status != "pending":
            return
        export.status = "running"
        db.commit()
        try:
            EXPORT_ROOT.mkdir(parents=True, exist_ok=True)
            suffix = "xlsx" if export.export_type == "excel" else "csv"
            path = EXPORT_ROOT / f"{export.id}.{suffix}"
            if export.export_type == "excel":
                _write_excel(db, export.website_id, path)
            else:
                _write_csv(db, export, path)
            export.file_path = str(path)
            export.status = "succeeded"
            export.finished_at = datetime.now(UTC)
            db.commit()
        except Exception as exc:
            db.rollback()
            export = db.get(Export, uuid.UUID(export_id))
            if export:
                export.status = "failed"
                export.error_message = str(exc)[:4000]
                export.finished_at = datetime.now(UTC)
                db.commit()
            raise


def _datasets(
    db: Session,
    website_id: object,
    *,
    selected_type: str | None = None,
    item_ids: list[str] | None = None,
) -> dict[str, tuple[list[str], list[list[object]]]]:
    selected_ids = {uuid.UUID(item_id) for item_id in item_ids or []}
    urls = list(db.scalars(select(Url).where(Url.website_id == website_id)))
    url_ids = [url.id for url in urls]
    url_by_id = {url.id: url.normalized_url for url in urls}
    issues = list(db.scalars(select(Issue).where(Issue.website_id == website_id)))
    changes = list(db.scalars(select(Change).where(Change.website_id == website_id)))
    job_listings = list(
        db.scalars(
            select(JobListing)
            .where(JobListing.website_id == website_id)
            .order_by(JobListing.valid_through.asc().nullslast(), JobListing.title)
        )
    )
    tasks = list(
        db.scalars(
            select(RecommendationTask)
            .where(RecommendationTask.website_id == website_id)
            .order_by(RecommendationTask.updated_at.desc())
        )
    )
    if item_ids is not None:
        if selected_type == "urls":
            urls = [url for url in urls if url.id in selected_ids]
            url_ids = [url.id for url in urls]
            url_by_id = {url.id: url.normalized_url for url in urls}
        elif selected_type == "changes":
            changes = [change for change in changes if change.id in selected_ids]
        elif selected_type == "vacancies":
            job_listings = [listing for listing in job_listings if listing.id in selected_ids]
        elif selected_type == "tasks":
            tasks = [task for task in tasks if task.id in selected_ids]
    source_rows = (
        list(db.scalars(select(UrlSource).where(UrlSource.url_id.in_(url_ids))))
        if url_ids
        else []
    )
    sources_by_url: dict[object, list[UrlSource]] = {}
    for source in source_rows:
        sources_by_url.setdefault(source.url_id, []).append(source)
    latest_full_run = db.scalar(
        select(CrawlRun)
        .where(CrawlRun.website_id == website_id, CrawlRun.crawl_type == "full_site_crawl")
        .order_by(CrawlRun.started_at.desc())
        .limit(1)
    )
    assigned_user_ids = {
        task.assigned_to_user_id for task in tasks if task.assigned_to_user_id
    }
    users = {
        user.id: user
        for user in db.scalars(select(User).where(User.id.in_(assigned_user_ids)))
    }
    issues_by_id = {issue.id: issue for issue in issues}
    job_issues_by_url = {
        listing.url_id: [
            issue
            for issue in issues
            if issue.url_id == listing.url_id
            and issue.issue_type in JOB_ISSUE_TYPES
            and issue.status in ACTIVE_JOB_ISSUE_STATUSES
        ]
        for listing in job_listings
    }
    latest_full_run_id = db.scalar(
        select(CrawlRun.id)
        .where(
            CrawlRun.website_id == website_id,
            CrawlRun.crawl_type == "full_site_crawl",
            CrawlRun.status == "succeeded",
        )
        .order_by(CrawlRun.finished_at.desc())
        .limit(1)
    )
    links = (
        list(
            db.execute(
                select(
                    UrlLink.source_url_id,
                    UrlLink.target_url_id,
                    UrlLink.target_url,
                    UrlLink.anchor_text,
                    UrlLink.is_internal,
                    UrlLink.is_nofollow,
                    UrlLink.http_status,
                )
                .distinct()
                .where(
                    UrlLink.source_url_id.in_(url_ids),
                    UrlLink.crawl_run_id == latest_full_run_id,
                )
            )
        )
        if url_ids and latest_full_run_id
        else []
    )
    return {
        "urls": (
            [
                "url",
                "status_code",
                "is_active",
                "is_indexable",
                "first_seen_at",
                "last_seen_at",
                "final_url",
                "page_type",
                "crawl_depth",
                "all_sources",
                "current_sources",
                "historical_sources",
                "last_light_checked_at",
                "last_full_analyzed_at",
            ],
            [
                [
                    url.normalized_url,
                    url.current_status_code,
                    url.is_active,
                    url.is_indexable,
                    url.first_seen_at,
                    url.last_seen_at,
                    url.current_final_url,
                    url.page_type,
                    url.crawl_depth,
                    " | ".join(
                        sorted(
                            {source.source_type for source in sources_by_url.get(url.id, [])}
                        )
                    ),
                    " | ".join(
                        _current_source_types(sources_by_url.get(url.id, []), latest_full_run)
                    ),
                    " | ".join(
                        _historical_source_types(sources_by_url.get(url.id, []), latest_full_run)
                    ),
                    url.last_light_checked_at,
                    url.last_full_analyzed_at,
                ]
                for url in urls
            ],
        ),
        "issues": (
            [
                "url",
                "type",
                "category",
                "severity",
                "status",
                "title",
                "first_detected",
                "last_detected",
            ],
            [
                [
                    url_by_id.get(issue.url_id),
                    issue.issue_type,
                    issue.category,
                    issue.severity,
                    issue.status,
                    issue.title,
                    issue.first_detected_at,
                    issue.last_detected_at,
                ]
                for issue in issues
            ],
        ),
        "changes": (
            ["url", "type", "field", "old_value", "new_value", "detected_at"],
            [
                [
                    url_by_id.get(change.url_id),
                    change.change_type,
                    change.field_name,
                    change.old_value,
                    change.new_value,
                    change.detected_at,
                ]
                for change in changes
            ],
        ),
        "links": (
            [
                "source_url",
                "target_url",
                "anchor_text",
                "internal",
                "nofollow",
                "status",
            ],
            [
                [
                    url_by_id.get(link.source_url_id),
                    link.target_url,
                    link.anchor_text,
                    link.is_internal,
                    link.is_nofollow,
                    link.http_status,
                ]
                for link in links
            ],
        ),
        "vacancies": (
            [
                "url",
                "title",
                "employer",
                "locations",
                "lifecycle_status",
                "google_for_jobs_status",
                "date_posted",
                "valid_through",
                "employment_types",
                "hours",
                "salary",
                "application_url",
                "status_code",
                "is_indexable",
                "inbound_internal_links",
                "active_findings",
                "first_detected_at",
                "last_detected_at",
            ],
            [
                [
                    url_by_id.get(listing.url_id),
                    listing.title,
                    listing.employer,
                    " | ".join(listing.locations or []),
                    listing.lifecycle_status,
                    _job_validation_status(
                        listing,
                        job_issues_by_url.get(listing.url_id, []),
                    ),
                    listing.date_posted,
                    listing.valid_through,
                    " | ".join(listing.employment_types or []),
                    listing.hours,
                    json.dumps(listing.salary_data, ensure_ascii=False)
                    if listing.salary_data
                    else "",
                    listing.application_url,
                    listing.current_status_code,
                    listing.is_indexable,
                    listing.inbound_internal_links,
                    " | ".join(issue.title for issue in job_issues_by_url.get(listing.url_id, [])),
                    listing.first_detected_at,
                    listing.last_detected_at,
                ]
                for listing in job_listings
            ],
        ),
        "tasks": (
            [
                "title",
                "action",
                "issue_url",
                "category",
                "priority",
                "priority_reason",
                "status",
                "primary_role",
                "supporting_roles",
                "assigned_to",
                "effort_min_minutes",
                "effort_max_minutes",
                "verification_status",
                "required_input",
                "acceptance_criteria",
                "created_at",
                "updated_at",
            ],
            [
                [
                    task.title,
                    task.action,
                    url_by_id.get(issues_by_id[task.primary_issue_id].url_id)
                    if task.primary_issue_id in issues_by_id
                    else None,
                    task.category,
                    task.priority,
                    task.priority_reason,
                    task.status,
                    task.primary_role,
                    " | ".join(task.supporting_roles or []),
                    _user_label(users.get(task.assigned_to_user_id)),
                    task.effort_min_minutes,
                    task.effort_max_minutes,
                    task.verification_status,
                    " | ".join(task.required_input or []),
                    " | ".join(task.acceptance_criteria or []),
                    task.created_at,
                    task.updated_at,
                ]
                for task in tasks
            ],
        ),
    }


def _current_source_types(sources: list[UrlSource], run: CrawlRun | None) -> list[str]:
    return sorted(
        {
            source.source_type
            for source in sources
            if source.source_type == "manual"
            or run is None
            or _seen_since(source.last_seen_at, run.started_at)
        }
    )


def _historical_source_types(sources: list[UrlSource], run: CrawlRun | None) -> list[str]:
    current = set(_current_source_types(sources, run))
    return sorted({source.source_type for source in sources} - current)


def _seen_since(seen_at: datetime, started_at: datetime) -> bool:
    return seen_at.replace(tzinfo=None) >= started_at.replace(tzinfo=None)


def _user_label(user: User | None) -> str:
    if user is None:
        return ""
    return user.display_name or user.email


def _job_validation_status(listing: JobListing, issues: list[Issue]) -> str:
    if any(issue.severity in {"critical", "high"} for issue in issues):
        return "error"
    if issues:
        return "warning"
    if "job_posting_schema" in (listing.detection_sources or []):
        return "valid"
    return "missing_schema"


def _write_csv(db: Session, export: Export, path: Path) -> None:
    datasets = _datasets(
        db,
        export.website_id,
        selected_type=export.export_type,
        item_ids=export.item_ids,
    )
    key = "issues" if export.export_type == "technical" else export.export_type
    if key not in datasets:
        raise ValueError("Unsupported CSV export type")
    headers, rows = datasets[key]
    website = db.get(Website, export.website_id)
    exported_at = datetime.now(UTC).isoformat()
    filter_text = json.dumps(export.filters or {}, ensure_ascii=False, sort_keys=True)
    metadata = [website.name if website else str(export.website_id), exported_at, filter_text]
    headers = [*headers, "website", "exported_at", "filters"]
    rows = [[*row, *metadata] for row in rows]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def _write_excel(db: Session, website_id: object, path: Path) -> None:
    website = db.get(Website, website_id)
    workbook = Workbook()
    metadata = workbook.active
    metadata.title = "Metadata"
    metadata.append(["Website", website.name if website else str(website_id)])
    metadata.append(["Exported at", datetime.now(UTC).isoformat()])
    for name, (headers, rows) in _datasets(db, website_id).items():
        sheet = workbook.create_sheet(name.title())
        sheet.append(headers)
        for row in rows:
            sheet.append([_excel_value(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if rows:
            table = Table(displayName=f"Table{name.title()}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    workbook.save(path)


def _excel_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.astimezone(UTC).replace(tzinfo=None)
    if isinstance(value, uuid.UUID):
        return str(value)
    return value
